#!/usr/bin/env python3
"""Build the blank, scenario-based student IPS. Original long framework is preserved."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'files/BUS331_Abbreviated_IPS_Template.docx'

def main():
    d = Document()
    s = d.sections[0]
    s.top_margin = s.bottom_margin = Inches(.6)
    s.left_margin = s.right_margin = Inches(.7)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        st = d.styles[name]
        st.font.name = 'Calibri'
        st.font.color.rgb = RGBColor(0, 0, 0)
        ppr = st.element.find(qn('w:pPr'))
        if ppr is not None:
            for border in list(ppr.findall(qn('w:pBdr'))): ppr.remove(border)
    d.styles['Normal'].font.size = Pt(10)
    d.styles['Normal'].paragraph_format.space_after = Pt(5)
    d.styles['Title'].font.size = Pt(21)
    d.styles['Heading 1'].font.size = Pt(13)
    d.styles['Heading 1'].paragraph_format.space_before = Pt(8)
    d.styles['Heading 1'].paragraph_format.space_after = Pt(4)
    def p(text): return d.add_paragraph(text)
    def h(text): d.add_heading(text, level=1)
    def field(text): p(text + '\n[Enter a concise response]')
    def table(headers, rows):
        t=d.add_table(rows=1, cols=len(headers)); t.autofit=False
        for c, x in zip(t.rows[0].cells, headers):
            c.text=x
            fill=OxmlElement('w:shd'); fill.set(qn('w:fill'),'E8EDF2'); c._tc.get_or_add_tcPr().append(fill)
            for r in c.paragraphs[0].runs: r.bold=True
        for row in rows:
            for c,x in zip(t.add_row().cells,row): c.text=x
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
        t._tbl.tblPr.append(borders)
        for row in t.rows:
            keep=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(keep)
            for c in row.cells:
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(4)
        return t
    d.add_heading('Investment Policy Statement',0)
    p('BUS331 Investments | Phase 1 client mandate | One copy per assigned client')
    p('Use the assigned profile slides and client-data row. Cite the slide and row for material facts. Keep responses brief; bullets are welcome. The blank form is two pages and may expand as you type.')
    p('Missing client information: write “Not provided” and state the investment implication only if material. Do not invent account balances, expenses, tax rates, contribution room, insurance, or family details. Label interpretations and any essential modeling assumption; proposed policies are committee recommendations, not client facts.')
    p('Client and case number: ____________________  Team: ____________________\nProfile slide and workbook row: __________________________________________')
    h('1 Client goals and supplied facts')
    field('Priority goal and any competing goal; include stated amount and timing only when supplied')
    field('Relevant background and constraint evidence from the case; net worth is not automatically investable assets')
    h('2 Return and risk mandate')
    table(['Case inputs from the client data', 'Committee proposed policy and brief rationale'], [
        ('Target annual E(R): _____%','Annual return objective: _____%\n[Explain feasibility using the submitted macro view; no guaranteed return]'),
        ('Standard deviation σ: _____%','Annual portfolio volatility limit: _____%\n[Explain whether you retain or revise the case input]'),
        ('Risk classification: __________\nRisk aversion A: _____','Drawdown review tripwire: _____% from peak\n[Justify from the case; this is a review trigger, not a loss guarantee]')])
    field('Risk capacity versus willingness: one sentence on each, and the main conflict or trade-off')
    p('Copy the approved objective, volatility limit, and drawdown tripwire into the Decision Record. Volatility and drawdown measure different risks; do not treat σ as a maximum loss.')
    d.add_page_break()
    h('3 Constraints that affect the portfolio')
    p('Together with Section 2, this covers RRTTLLU: Risk, Return, Time horizon, Tax, Liquidity, Legal, and Unique circumstances. “Not provided” is a valid response for missing facts; it does not mean no constraint exists.')
    table(['Constraint', 'Supplied fact and portfolio implication'], [
        ('Time horizon','[State supplied timing; distinguish near-term goals from long-term wealth]'),
        ('Liquidity','[State supplied cash need and date or reserve requirement; leave an unknown amount unknown]'),
        ('Tax','[State any supplied tax concern; otherwise Not provided]'),
        ('Legal and unique','[State supplied restrictions, ethical preferences, concentration or special circumstances; otherwise Not provided]')])
    h('4 Provisional strategic allocation')
    p('Recommend broad policy weights totaling 100%. These are committee choices supported by the case and submitted macro view. Specific securities and final optimized weights belong to Phase 2.')
    table(['Equity', 'Fixed income', 'Cash and equivalents', 'Other if justified', 'Total'], [('_____%','_____%','_____%','_____%','100%')])
    field('Allocation rationale and principal trade-off; explain protection of any stated near-term cash need')
    p('If investable assets are not supplied, work in percentages and flag whether the cash need can be funded as unresolved. Do not calculate a cash weight from total net worth without an explicit, justified modeling assumption.')
    h('5 Implementation and review policy')
    p('Use funds and ETFs as the primary vehicles. Phase 2 selects 8–10 holdings, never more than 10, with no more than two or three individual securities. Apply every stated client restriction.')
    field('Committee review and rebalancing rule: when to check weights, return, volatility, liquidity and restrictions; what triggers action')
    p('If the drawdown tripwire or another mandate constraint is breached, document the cause, proposed correction and full re-test before approval. Evaluate any hedge or no-hedge decision in Phase 2.')
    h('6 Open issues and committee approval')
    field('Material unresolved fact or essential modeling assumption, its basis and what decision could change; write None if none')
    p('Gate 1 status: ☐ Approve  ☐ Revise  ☐ Reject\nDecision Record vote reference and required action: __________________________')
    p('No client interview or signature is required. Combine the three completed IPS documents in assigned-client order as BUS331_[TeamName]_Phase1_ClientIPS.pdf.')
    footer=s.footer.paragraphs[0]
    footer.text='BUS331 | IPS | '
    page=OxmlElement('w:fldSimple'); page.set(qn('w:instr'),'PAGE'); footer._p.append(page)
    d.core_properties.title='BUS331 Investment Policy Statement'
    d.core_properties.author='Professor Evitts'
    d.save(OUTPUT)
    print(OUTPUT)

if __name__ == '__main__': main()
